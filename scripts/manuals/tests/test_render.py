"""A failed test or changed screenshot must never publish a verified user manual."""
import hashlib
import json
from pathlib import Path

import pytest
from docx import Document
from PIL import Image

from scripts.manuals.recording import DISCLAIMER
from scripts.manuals.render import build_manuals


@pytest.fixture
def evidence(tmp_path):
    folder = tmp_path / "evidence" / "sample"
    folder.mkdir(parents=True)
    image = folder / "01.png"
    Image.new("RGB", (32, 24), "white").save(image)
    manifest = {
        "schema_version": 1, "scenario": "sample", "title": "Sample <claim> workflow", "status": "passed",
        "purpose": "Exercise a sample UI", "scope": "One observed step", "prerequisites": ["Synthetic stack"],
        "limitations": ["Not an official manual"], "disclaimer": DISCLAIMER,
        "finished_at": "2026-09-30T12:00:00+00:00", "checkout": {"commit": "sample-commit", "dirty": True},
        "steps": [{"role": "Member", "number": 1, "title": "Check <claim>", "instruction": "Select Claim.",
                   "expected": "Claim form is visible.", "url": "http://localhost:5173/member/claims",
                   "captured_at": "2026-09-30T12:00:00+00:00", "screenshot": "01.png",
                   "sha256": hashlib.sha256(image.read_bytes()).hexdigest()}],
    }
    path = folder / "manifest.json"
    path.write_text(json.dumps(manifest))
    return tmp_path, path, manifest


@pytest.mark.parametrize("status", ["running", "failed", "skipped"])
def test_unverified_capture_cannot_publish(evidence, status):
    output, path, manifest = evidence
    manifest["status"] = status
    path.write_text(json.dumps(manifest))
    with pytest.raises(ValueError, match="incomplete or failed"):
        build_manuals(output, pdf=False)
    assert not (output / "manuals").exists()


def test_changed_screenshot_cannot_publish(evidence):
    output, path, _ = evidence
    (path.parent / "01.png").write_bytes(b"changed evidence")
    with pytest.raises(ValueError, match="Screenshot evidence changed"):
        build_manuals(output, pdf=False)
    assert not (output / "manuals").exists()


def test_screenshot_cannot_escape_evidence_directory(evidence):
    output, path, manifest = evidence
    manifest["steps"][0]["screenshot"] = "../01.png"
    path.write_text(json.dumps(manifest))
    with pytest.raises(ValueError, match="inside this scenario"):
        build_manuals(output, pdf=False)


def test_verified_evidence_produces_role_manual_and_embedded_screenshot(evidence):
    output, _, _ = evidence
    index = build_manuals(output, pdf=False)
    assert index.is_file()
    html = (index.parent / "sample-member.html").read_text()
    assert "data:image/png;base64," in html
    assert "Check &lt;claim&gt;" in html
    assert "Expected result:" in html
    assert DISCLAIMER in html
    document = Document(index.parent / "sample-member.docx")
    assert len(document.inline_shapes) == 1
    assert any("Claim form is visible" in p.text for p in document.paragraphs)
    assert len(document.tables) == 1
    catalogue = json.loads((index.parent / "catalogue.json").read_text())
    assert catalogue["status"] == "verified"
    assert catalogue["manuals"][0]["role"] == "Member"


def test_empty_capture_cannot_publish(tmp_path):
    with pytest.raises(ValueError, match="No UI evidence"):
        build_manuals(tmp_path, pdf=False)
