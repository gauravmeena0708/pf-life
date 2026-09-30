"""Render role-specific DOCX, HTML and PDF manuals from complete, verified UI evidence."""
import base64
import hashlib
import html
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from scripts.manuals.recording import DISCLAIMER

PROCESS_ROLE = "Cross-role process"


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def load_evidence(output: Path) -> list[tuple[dict, Path]]:
    manifests = sorted((output / "evidence").glob("*/manifest.json"))
    if not manifests:
        raise ValueError("No UI evidence found; run the browser tests first")
    verified = []
    for path in manifests:
        manifest = json.loads(path.read_text(encoding="utf-8"))
        if manifest.get("schema_version") != 1 or manifest.get("status") != "passed" or not manifest.get("steps"):
            raise ValueError(f"Cannot publish an incomplete or failed scenario: {path.parent.name}")
        if manifest.get("disclaimer") != DISCLAIMER:
            raise ValueError("The evidence must identify the synthetic POC")
        lifecycle = manifest.get("lifecycle")
        if lifecycle and (not lifecycle.get("expected_outcome") or lifecycle.get("outcome") != lifecycle["expected_outcome"]
                          or not lifecycle.get("observations")):
            raise ValueError("The lifecycle must reach its asserted outcome before publication")
        for step in manifest["steps"]:
            image = (path.parent / step["screenshot"]).resolve()
            if image.parent != path.parent.resolve() or image.suffix != ".png":
                raise ValueError("Screenshot must be a PNG inside this scenario's evidence directory")
            if hashlib.sha256(image.read_bytes()).hexdigest() != step["sha256"]:
                raise ValueError(f"Screenshot evidence changed: {image.name}")
        verified.append((manifest, path.parent))
    return verified


def metadata(manifest: dict, role: str) -> list[tuple[str, str]]:
    rows = [
        ("Manual version", "1.0 — generated from a verified UI run"),
        ("Role", role), ("Captured (UTC)", manifest["finished_at"]),
        ("Checkout commit", manifest["checkout"]["commit"]),
        ("Uncommitted checkout changes", "Yes" if manifest["checkout"]["dirty"] else "No"),
        ("Runtime provenance", "Screenshots and assertions establish the observed UI. The running server may differ from the checkout."),
    ]
    if manifest.get("runtime"):
        rows += [("Isolated runtime fixture", manifest["runtime"]["project"]),
                 ("Runtime image provenance", "Exact running image IDs are recorded in runtime.json alongside these manuals.")]
    return rows


def lifecycle_facts(manifest):
    lifecycle = manifest.get("lifecycle")
    if not lifecycle:
        return []
    return [("Case IDs", ", ".join(lifecycle["case_ids"])), ("Input data", lifecycle["data"]),
            ("Expected outcome", lifecycle["expected_outcome"]), ("Asserted outcome", lifecycle["outcome"]),
            *[(item["stage"], f"{item['value']} (evidence step {item['step']})") for item in lifecycle["observations"]]]


def write_html(manifest: dict, folder: Path, role: str, steps: list[dict], destination: Path) -> None:
    esc = html.escape
    css = Path(__file__).with_name("manual.css").read_text(encoding="utf-8")
    parts = [f'<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
             f'<title>{esc(manifest["title"])} — {esc(role)}</title><style>{css}</style></head><body>',
             f'<header><p>PF Life · User manual</p><h1>{esc(manifest["title"])}</h1><p class="role">Role: {esc(role)}</p></header>',
             f'<main><p class="notice">{esc(DISCLAIMER)}</p><section><h2>Version control chart</h2><table><tbody>']
    parts.extend(f'<tr><th scope="row">{esc(key)}</th><td>{esc(value)}</td></tr>' for key, value in metadata(manifest, role))
    parts += ['</tbody></table></section><nav aria-label="Contents"><h2>Contents</h2><ol>',
              '<li><a href="#introduction">Introduction</a></li>']
    parts.extend(f'<li><a href="#step-{i}">{esc(step["title"])}</a></li>' for i, step in enumerate(steps, 1))
    parts += ['<li><a href="#limits">Coverage and disclaimer</a></li></ol></nav>',
              f'<section id="introduction"><h2>Introduction</h2><h3>Purpose</h3><p>{esc(manifest["purpose"])}</p>'
              f'<h3>Scope</h3><p>{esc(manifest["scope"])}</p><p>{"This process manual follows all roles in execution order." if role == PROCESS_ROLE else "This role manual contains only the steps observed for " + esc(role) + " in this run."}</p>'
              '<h3>Prerequisites</h3><ul>']
    parts.extend(f'<li>{esc(item)}</li>' for item in manifest["prerequisites"])
    parts += ['</ul></section>']
    if lifecycle_facts(manifest):
        parts += ['<section><h2>Lifecycle data and asserted outcomes</h2><table><tbody>']
        parts.extend(f'<tr><th scope="row">{esc(k)}</th><td>{esc(v)}</td></tr>' for k, v in lifecycle_facts(manifest))
        parts += ['</tbody></table></section>']
    for i, step in enumerate(steps, 1):
        encoded = base64.b64encode((folder / step["screenshot"]).read_bytes()).decode("ascii")
        parts.append(f'<section class="step" id="step-{i}"><h2>{i}. {esc(step["title"])}</h2><p><strong>Role:</strong> {esc(step["role"])}</p><p>{esc(step["instruction"])}</p>'
                     f'<p class="expected"><strong>Expected result:</strong> {esc(step["expected"])}</p>'
                     f'<figure><img src="data:image/png;base64,{encoded}" alt="{esc(step["title"])} — observed POC screen">'
                     f'<figcaption>Figure {i}. {esc(step["title"])} · {esc(step["captured_at"])}</figcaption></figure>'
                     f'<p class="evidence">Screen: {esc(step["url"])}<br>Screenshot SHA-256: {esc(step["sha256"])}</p></section>')
    parts += ['<section id="limits"><h2>Coverage and disclaimer</h2><ul>']
    parts.extend(f'<li>{esc(item)}</li>' for item in manifest["limitations"])
    parts += [f'</ul><p>{esc(DISCLAIMER)}. These instructions describe the captured POC workflow. They are not official EPFO manuals '
              'and do not override statutory instructions, circulars or official guidance. Passwords and one-time codes are masked in screenshots.</p>'
              '</section></main><footer>PF Life · Generated from a successful browser test</footer></body></html>']
    destination.write_text("\n".join(parts), encoding="utf-8")


def write_docx(manifest: dict, folder: Path, role: str, steps: list[dict], destination: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.8)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Trebuchet MS", Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    for name, size in (("Title", 26), ("Heading 1", 17), ("Heading 2", 13)):
        style = document.styles[name]
        style.font.name, style.font.size = "Georgia", Pt(size)
        style.font.color.rgb = RGBColor.from_string("003366")
    section.header.paragraphs[0].text = "PF Life | Synthetic POC user manual"
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Synthetic demonstration | Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = f'{manifest["title"]} — {role}'
    document.core_properties.subject = DISCLAIMER
    document.core_properties.author = "PF Life UI test automation"
    document.add_paragraph("PF LIFE · USER MANUAL", "Subtitle")
    document.add_heading(manifest["title"], 0)
    document.add_paragraph(f"Role: {role}", "Subtitle")
    banner = document.add_paragraph(DISCLAIMER)
    banner.runs[0].bold = True
    banner.runs[0].font.color.rgb = RGBColor.from_string("B92125")
    document.add_heading("Version control chart", 1)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in metadata(manifest, role):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    document.add_page_break()
    document.add_heading("Contents", 1)
    document.add_paragraph("Introduction")
    for i, step in enumerate(steps, 1):
        document.add_paragraph(f'{i}. {step["title"]}')
    document.add_paragraph("Coverage and disclaimer")
    document.add_heading("Introduction", 1)
    document.add_heading("Purpose", 2)
    document.add_paragraph(manifest["purpose"])
    document.add_heading("Scope", 2)
    document.add_paragraph(manifest["scope"])
    document.add_paragraph("This process manual follows all roles in execution order." if role == PROCESS_ROLE
                           else f"This role manual contains only the steps observed for {role} in this run.")
    document.add_heading("Prerequisites", 2)
    for item in manifest["prerequisites"]:
        document.add_paragraph(item, "List Bullet")
    if lifecycle_facts(manifest):
        document.add_heading("Lifecycle data and asserted outcomes", 1)
        facts = document.add_table(rows=0, cols=2)
        facts.style = "Table Grid"
        for label, value in lifecycle_facts(manifest):
            cells = facts.add_row().cells
            cells[0].text, cells[1].text = label, value
    for i, step in enumerate(steps, 1):
        document.add_page_break()
        document.add_heading(f'{i}. {step["title"]}', 1)
        document.add_paragraph("Role: " + step["role"])
        document.add_paragraph(step["instruction"])
        document.add_paragraph("Expected result: " + step["expected"])
        document.add_picture(str(folder / step["screenshot"]), width=Inches(6.6))
        document.add_paragraph(f'Figure {i}. {step["title"]}', "Caption")
        document.add_paragraph(f'Captured: {step["captured_at"]}\nScreen: {step["url"]}')
        document.add_paragraph("Screenshot SHA-256: " + step["sha256"])
    document.add_page_break()
    document.add_heading("Coverage and disclaimer", 1)
    for item in manifest["limitations"]:
        document.add_paragraph(item, "List Bullet")
    document.add_paragraph(DISCLAIMER + ". These instructions describe the captured POC workflow. They are not official EPFO manuals "
                           "and do not override statutory instructions, circulars or official guidance. "
                           "Passwords and one-time codes are masked in screenshots.")
    document.save(destination)


def build_manuals(output: Path, pdf: bool = True) -> Path:
    # Validate every scenario before writing anything. Failed capture must not produce a verified manual.
    evidence = load_evidence(output)
    folder = output / "manuals"
    folder.mkdir(exist_ok=False)
    entries = []
    for manifest, images in evidence:
        roles = list(dict.fromkeys(step["role"] for step in manifest["steps"]))
        if manifest.get("lifecycle"):
            roles.append(PROCESS_ROLE)
        for role in roles:
            steps = manifest["steps"] if role == PROCESS_ROLE else [step for step in manifest["steps"] if step["role"] == role]
            name = slug(manifest["scenario"] + "-" + role)
            write_html(manifest, images, role, steps, folder / f"{name}.html")
            write_docx(manifest, images, role, steps, folder / f"{name}.docx")
            entries.append({"name": name, "title": manifest["title"], "role": role, "steps": len(steps),
                            "kind": "process" if role == PROCESS_ROLE else "role",
                            "case_ids": manifest.get("lifecycle", {}).get("case_ids", [])})
    if pdf:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            browser = p.chromium.launch()
            try:
                page = browser.new_page()
                for entry in entries:
                    page.goto((folder / (entry["name"] + ".html")).resolve().as_uri())
                    page.pdf(path=str(folder / (entry["name"] + ".pdf")), format="A4", print_background=True,
                             margin={"top": "12mm", "bottom": "12mm", "left": "12mm", "right": "12mm"})
            finally:
                browser.close()
    rows = []
    for entry in entries:
        links = " · ".join(f'<a href="{entry["name"]}.{kind}">{kind.upper()}</a>' for kind in (["docx", "html", "pdf"] if pdf else ["docx", "html"]))
        rows.append(f'<tr><td>{html.escape(entry["title"])}</td><td>{html.escape(entry["role"])}</td><td>{entry["steps"]}</td><td>{links}</td></tr>')
    css = Path(__file__).with_name("manual.css").read_text(encoding="utf-8")
    (folder / "index.html").write_text(
        f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>PF Life user manuals</title><style>{css}</style></head>'
        f'<body><header><h1>PF Life user manuals</h1></header><main><p class="notice">{html.escape(DISCLAIMER)}</p>'
        '<p>Generated from verified browser journeys. Coverage is limited to the steps captured in each role manual.</p>'
        '<table><thead><tr><th>Workflow</th><th>Role</th><th>Steps</th><th>Downloads</th></tr></thead><tbody>'
        + "\n".join(rows) + '</tbody></table></main></body></html>', encoding="utf-8")
    (folder / "catalogue.json").write_text(json.dumps({"status": "verified", "manuals": entries}, indent=2) + "\n", encoding="utf-8")
    return folder / "index.html"
